"""Export Excel des tableaux d'analyse (v1.3.29).

Convertit les tableaux Markdown des livrables et des productions d'etapes en
un classeur .xlsx (un onglet par tableau) avec en-tete stylé, filtres
automatiques, volets figés et coloration des niveaux de risque/statut.

Le .md reste le document maitre (texte, justifications) — l'Excel est une vue
tabulaire produite en plus, sans modification du contenu.
"""

import re
from pathlib import Path

from config import config as app_config

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    OPENPYXL_OK = True
except Exception:  # pragma: no cover - openpyxl absent
    OPENPYXL_OK = False

_HEADER_FILL = "4472C4"
_SEV_FILL = {
    "critique": "C00000",
    "élevé": "ED7D31", "elevé": "ED7D31", "élevée": "ED7D31", "majeur": "ED7D31",
    "moyen": "FFC000", "modéré": "FFC000",
    "faible": "92D050", "mineur": "D9D9D9",
    "ko": "C00000", "ok": "92D050",
}

_SEP_RE = re.compile(r"^\s*\|?[\s:\-|]+\|?\s*$")


def _row_cells(line: str) -> list[str]:
    """Cellules d'une ligne de tableau markdown (tolere \\| echappe)."""
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    cells = re.split(r"(?<!\\)\|", s)
    out = []
    for c in cells:
        c = c.replace("\\|", "|")
        c = c.replace("<br>", "\n").replace("<br/>", "\n").replace("<br />", "\n")
        c = c.replace("**", "").replace("__", "")
        out.append(c.strip())
    return out


def _is_separator(line: str) -> bool:
    s = line.strip()
    if not s:
        return False
    return bool(re.match(r"^\s*\|?[\s:\-|]+\|?\s*$", s)) and "-" in s


def _sheet_name(title: str, used: set) -> str:
    name = re.sub(r"[\[\]:*?/\\]", " ", title)
    name = (re.sub(r"\s+", " ", name).strip() or "Feuille")[:31]
    base, k = name[:28], 2
    while name.lower() in {u.lower() for u in used}:
        name = f"{base}_{k}"[:31]
        k += 1
    used.add(name)
    return name


def markdown_tables(md: str) -> list[tuple[str, list[list[str]]]]:
    """Extrait les tableaux d'un markdown : [(titre, lignes avec en-tete)].

    Robustesse : en-tetes en gras, separateurs ':---', cellules <br>,
    lignes de continuation (repli du modele coupant une cellule longue)."""
    lines = (md or "").splitlines()
    tables: list[tuple[str, list[list[str]]]] = []
    current_title = ""
    i = 0
    n = len(lines)
    while i < n:
        stripped = lines[i].strip()
        m = re.match(r"^#{1,6}\s+(.*)$", stripped)
        if m:
            current_title = re.sub(r"[*_`]+", "", m.group(1)).strip()
        elif re.match(r"^\*\*.+\*\*:?\s*$", stripped):
            current_title = re.sub(r"[*_`]+", "", stripped).strip(": ")

        if stripped.startswith("|") and "|" in stripped[1:] and i + 1 < n \
                and _is_separator(lines[i + 1]):
            header = _row_cells(stripped)
            j = i + 2
            body: list[list[str]] = []
            while j < n:
                raw = lines[j]
                ls = raw.strip()
                if ls.startswith("|"):
                    body.append(_row_cells(ls))
                    j += 1
                    continue
                if ls == "":
                    # ligne vide : continuation si un tableau reprend apres
                    k = j
                    while k < n and not lines[k].strip():
                        k += 1
                    if k < n and lines[k].strip().startswith("|"):
                        j = k
                        continue
                    break
                if body:
                    # continuation de la derniere cellule (ligne sans pipe)
                    body[-1][-1] = f"{body[-1][-1]} {ls}".strip()
                    j += 1
                    continue
                break
            if header and body:
                tables.append((current_title or f"Tableau {len(tables) + 1}",
                               [header] + body))
            i = max(j, i + 1)
            current_title = ""
            continue
        i += 1
    return tables


def _write_sheet(wb, title: str, cells: list[list[str]], used: set) -> None:
    ws = wb.create_sheet(_sheet_name(title, used))
    ws.append([str(c) for c in (cells[0] if cells else [])])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _fill(_HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in cells[1:]:
        ws.append([str(c) for c in row])

    # largeurs de colonnes (premiere ligne de texte, bornees)
    ncols = max((len(r) for r in cells), default=1)
    for col in range(ncols):
        width = 12
        for r in cells[:60]:
            if col < len(r):
                first_line = (r[col] or "").split("\n")[0]
                width = max(width, min(len(first_line) + 2, 60))
        ws.column_dimensions[get_column_letter(col + 1)].width = width

    # retour a la ligne + coloration niveau/statut
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            key = (cell.value or "").strip().lower()
            if key in _SEV_FILL:
                cell.fill = _fill(_SEV_FILL[key])
                cell.font = Font(bold=True,
                                 color="FFFFFF" if key in ("critique", "ko") else "000000")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _fill(color: str):
    from openpyxl.styles import PatternFill
    return PatternFill("solid", fgColor=color)


def export_xlsx(sheets: list[tuple[str, list[list[str]]]], path: Path) -> Path:
    """Ecrit un classeur .xlsx : [(titre d'onglet, lignes)] ; retourne le chemin."""
    if not OPENPYXL_OK:
        raise RuntimeError("openpyxl indisponible (pip install openpyxl)")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    for title, cells in sheets:
        if not cells:
            continue
        _write_sheet(wb, title, cells, used)
    if not wb.sheetnames:
        ws = wb.create_sheet("Sans tableau")
        ws.append(["Aucun tableau detecte — le texte complet reste dans le .md"])
    wb.save(path)
    return path


def save_xlsx_from_outputs(outputs: dict, project: str, output_dir) -> Path:
    """Classeur d'analyse : onglets du livrable (5 blocs) puis des productions
    d'etapes. Appelé apres save_analysis_outputs (CLI et GUI)."""
    from agents.orchestrator import WORKFLOW_STEPS

    sheets: list[tuple[str, list[list[str]]]] = []
    liv = (outputs or {}).get("livraison", "")
    if liv:
        for title, cells in markdown_tables(liv):
            sheets.append((title, cells))
    for step in WORKFLOW_STEPS:
        if step["id"] == "livraison":
            continue
        text = (outputs or {}).get(step["id"], "")
        for k, (title, cells) in enumerate(markdown_tables(text), 1):
            name = f"{step['name'].split(' - ', 1)[-1]}"
            sheets.append((f"{name} · {title}" if title else name, cells))
    path = Path(output_dir) / f"{project}_analyse.xlsx"
    return export_xlsx(sheets, path)
